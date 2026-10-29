import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import app as pdf4um_app
from PIL import Image
from docx.shared import Pt, RGBColor
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app import create_app


def make_pdf(page_count: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=300, height=400)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def make_text_pdf(text: str) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=400)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    page[NameObject("/Resources")] = writer._add_object(resources)
    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(content)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def make_image(format_name: str, color: str, size: tuple[int, int] = (120, 80)) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", size, color).save(output, format=format_name)
    return output.getvalue()


def make_docx() -> bytes:
    output = io.BytesIO()
    document = pdf4um_app.Document()
    document.add_heading("Titolo documento", level=1)
    document.add_paragraph("Testo di prova")
    cells = document.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text = "Prima colonna"
    cells[1].text = "Seconda colonna"
    document.save(output)
    return output.getvalue()


def make_odt() -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr(
            "mimetype",
            "application/vnd.oasis.opendocument.text",
            compress_type=zipfile.ZIP_STORED,
        )
        archive.writestr(
            "content.xml",
            '''<?xml version="1.0" encoding="UTF-8"?>
            <office:document-content
              xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
              xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
              xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
              xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
              office:version="1.2">
              <office:body><office:text>
                <text:h text:outline-level="1">Titolo documento</text:h>
                <text:p>Testo <text:span>di prova</text:span></text:p>
                <text:list text:style-name="Numbered">
                  <text:list-item><text:p>Elemento elenco</text:p></text:list-item>
                </text:list>
                <table:table>
                  <table:table-header-rows><table:table-row>
                    <table:table-cell><text:p>Prima colonna</text:p></table:table-cell>
                    <table:table-cell><text:p>Seconda colonna</text:p></table:table-cell>
                  </table:table-row></table:table-header-rows>
                  <table:table-row-group><table:table-row>
                    <table:table-cell><text:p>Terza colonna</text:p></table:table-cell>
                    <table:table-cell><text:p>Quarta colonna</text:p></table:table-cell>
                  </table:table-row></table:table-row-group>
                </table:table>
              </office:text></office:body>
            </office:document-content>''',
        )
        archive.writestr(
            "styles.xml",
            '''<?xml version="1.0" encoding="UTF-8"?>
            <office:document-styles
              xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
              xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
              xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0">
              <office:styles>
                <text:list-style style:name="Numbered">
                  <text:list-level-style-number text:level="1" style:num-format="1"/>
                </text:list-style>
              </office:styles>
            </office:document-styles>''',
        )
    return output.getvalue()


class PdfWorkbenchTest(unittest.TestCase):
    def setUp(self):
        web_directory = Path(__file__).parents[1] / "src" / "main" / "resources" / "web"
        self.web_directory_patcher = patch.object(pdf4um_app, "WEB_DIRECTORY", web_directory)
        self.web_directory_patcher.start()
        self.addCleanup(self.web_directory_patcher.stop)
        self.client = create_app().test_client()
        self.pdf = make_pdf()

    @staticmethod
    def upload(name: str, content: bytes):
        return (io.BytesIO(content), name)

    def post(self, path: str, data: dict):
        response = self.client.post(
            path, data=data, content_type="multipart/form-data", buffered=True
        )
        self.addCleanup(response.close)
        return response

    def get(self, path: str):
        response = self.client.get(path, buffered=True)
        self.addCleanup(response.close)
        return response

    def zip_names(self, response):
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            return archive.namelist()

    def test_serves_workbench_and_security_headers(self):
        response = self.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"PDF", response.data)
        self.assertIn(b'href="/" aria-label="Torna alla home di In4matika"', response.data)
        self.assertIn(
            b'src="logoinfo.png" alt="In4matika"',
            response.data,
        )
        self.assertIn(b'class="home-link" href="/">Home</a>', response.data)
        self.assertNotIn(b"Chi Siamo", response.data)
        self.assertNotIn(b"Servizi", response.data)
        self.assertNotIn(b"Contattaci", response.data)
        self.assertIn(b'value="images-to-pdf"', response.data)
        self.assertIn(b"Converti immagini in PDF", response.data)
        self.assertIn(b'value="documents-to-pdf"', response.data)
        self.assertIn(b'value="docx-to-pdf"', response.data)
        self.assertIn(b"Converti DOCX in PDF", response.data)
        self.assertIn(b"Il documento viene elaborato sul server", response.data)
        self.assertNotIn(b"docx-preview@", response.data)
        self.assertNotIn(b"html2canvas@", response.data)
        self.assertNotIn(b"jspdf@", response.data)
        self.assertIn(b'href="styles.css"', response.data)
        self.assertIn(b'src="app.js"', response.data)
        self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")
        javascript = self.get("/app.js")
        self.assertEqual(javascript.status_code, 200)
        self.assertIn(b'fetch("api/"', javascript.data)
        self.assertEqual(self.get("/styles.css").status_code, 200)

    def test_serves_uploaded_logo_from_web_assets_directory(self):
        with tempfile.TemporaryDirectory() as directory:
            logo = Path(directory) / "logoinfo.png"
            logo.write_bytes(b"\x89PNG\r\n\x1a\n")
            with patch.object(pdf4um_app, "WEB_DIRECTORY", Path(directory)):
                response = self.get("/logoinfo.png")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.mimetype, "image/png")
            self.assertEqual(response.data, b"\x89PNG\r\n\x1a\n")

    def test_serves_flat_named_assets_without_overwriting_site_homepage(self):
        with tempfile.TemporaryDirectory() as directory:
            app_directory = Path(directory)
            (app_directory / "index.html").write_text("Main website")
            (app_directory / "pdf4um-index.html").write_text("PDF workbench")
            (app_directory / "pdf4um-app.js").write_text("workbench script")
            (app_directory / "pdf4um-styles.css").write_text("workbench styles")
            (app_directory / "pdf4um-logoinfo.png").write_bytes(b"\x89PNG\r\n\x1a\n")
            with (
                patch.object(pdf4um_app, "APP_DIRECTORY", app_directory),
                patch.object(pdf4um_app, "WEB_DIRECTORY", app_directory / "removed-folder"),
            ):
                page = self.get("/")
                javascript = self.get("/app.js")
                stylesheet = self.get("/styles.css")
                logo = self.get("/logoinfo.png")

            self.assertEqual(page.data, b"PDF workbench")
            self.assertEqual(javascript.data, b"workbench script")
            self.assertEqual(stylesheet.data, b"workbench styles")
            self.assertEqual(logo.data, b"\x89PNG\r\n\x1a\n")

    def test_serves_logo_uploaded_directly_beside_app(self):
        with tempfile.TemporaryDirectory() as directory:
            app_directory = Path(directory)
            (app_directory / "logoinfo.png").write_bytes(b"\x89PNG\r\n\x1a\n")
            with (
                patch.object(pdf4um_app, "APP_DIRECTORY", app_directory),
                patch.object(pdf4um_app, "WEB_DIRECTORY", app_directory / "removed-folder"),
            ):
                response = self.get("/logoinfo.png")

            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.mimetype, "image/png")
            self.assertEqual(response.data, b"\x89PNG\r\n\x1a\n")

    def test_prefers_deployed_web_assets_to_legacy_source_assets(self):
        with tempfile.TemporaryDirectory() as directory:
            app_directory = Path(directory)
            deployed_assets = app_directory / "pdf4um-web"
            legacy_assets = app_directory / "src" / "main" / "resources" / "web"
            for assets in (deployed_assets, legacy_assets):
                assets.mkdir(parents=True)
                for filename in ("index.html", "app.js", "styles.css"):
                    (assets / filename).touch()

            resolved = pdf4um_app._resolve_web_directory(app_directory, None)

            self.assertEqual(resolved, deployed_assets)

    def test_workbench_routes_work_when_mounted_under_wsgi_prefix(self):
        environ = {"SCRIPT_NAME": "/flask.wsgi"}
        page = self.client.get("/", environ_overrides=environ, buffered=True)
        self.addCleanup(page.close)
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'href="styles.css"', page.data)

        stylesheet = self.client.get("/styles.css", environ_overrides=environ, buffered=True)
        self.addCleanup(stylesheet.close)
        self.assertEqual(stylesheet.status_code, 200)

    def test_split_and_merge_preserve_page_counts_and_order(self):
        split = self.post(
            "/api/split",
            {"file": self.upload("Report.pdf", make_pdf(3)), "pagesPerPart": "2"},
        )
        self.assertEqual(split.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(split.data)) as archive:
            self.assertEqual(archive.namelist(), ["Report_0001.pdf", "Report_0002.pdf"])
            self.assertEqual(len(PdfReader(io.BytesIO(archive.read("Report_0001.pdf"))).pages), 2)
            self.assertEqual(len(PdfReader(io.BytesIO(archive.read("Report_0002.pdf"))).pages), 1)

        merged = self.post(
            "/api/merge",
            {
                "files": [
                    self.upload("First.pdf", self.pdf),
                    self.upload("Second.pdf", make_pdf(2)),
                ]
            },
        )
        self.assertEqual(merged.status_code, 200)
        self.assertIn("First_merged.pdf", merged.headers["Content-Disposition"])
        self.assertEqual(len(PdfReader(io.BytesIO(merged.data)).pages), 3)

    def test_converts_multiple_jpg_and_png_images_to_a4_pdf_in_order(self):
        response = self.post(
            "/api/images-to-pdf",
            {
                "images": [
                    self.upload("first.jpg", make_image("JPEG", "red")),
                    self.upload("second.png", make_image("PNG", "blue")),
                ]
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/pdf")
        self.assertIn("immagini.pdf", response.headers["Content-Disposition"])
        reader = PdfReader(io.BytesIO(response.data))
        self.assertEqual(len(reader.pages), 2)
        for index, page in enumerate(reader.pages):
            self.assertAlmostEqual(float(page.mediabox.width), 595.27, places=1)
            self.assertAlmostEqual(float(page.mediabox.height), 841.89, places=1)
            self.assertTrue(page.images)
            pixel = page.images[0].image.convert("RGB").getpixel((0, 0))
            if index == 0:
                self.assertGreater(pixel[0], pixel[2])
            else:
                self.assertGreater(pixel[2], pixel[0])

    def test_image_conversion_rejects_unsupported_and_unreadable_files(self):
        unsupported = self.post(
            "/api/images-to-pdf",
            {"images": self.upload("animation.gif", b"not an image")},
        )
        self.assertEqual(unsupported.status_code, 415)

        unreadable = self.post(
            "/api/images-to-pdf",
            {"images": self.upload("broken.png", b"not a PNG")},
        )
        self.assertEqual(unreadable.status_code, 422)

    def test_converts_docx_and_odt_documents_to_pdf(self):
        for filename, content in (
            ("Report.docx", make_docx()),
            ("Report.odt", make_odt()),
        ):
            with self.subTest(filename=filename):
                response = self.post(
                    "/api/documents-to-pdf",
                    {"file": self.upload(filename, content)},
                )

                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.mimetype, "application/pdf")
                self.assertIn("Report.pdf", response.headers["Content-Disposition"])
                reader = PdfReader(io.BytesIO(response.data))
                self.assertGreaterEqual(len(reader.pages), 1)
                extracted_text = "\n".join(page.extract_text() or "" for page in reader.pages)
                self.assertIn("Titolo documento", extracted_text)
                self.assertIn("Testo di prova", extracted_text)
                self.assertIn("Prima colonna", extracted_text)
                self.assertIn("Seconda colonna", extracted_text)
                if filename.endswith(".odt"):
                    self.assertIn("Elemento elenco", extracted_text)

    def test_docx_conversion_preserves_run_font_sizes_and_emphasis(self):
        document = pdf4um_app.Document()
        title = document.add_paragraph()
        title_run = title.add_run("Titolo personalizzato")
        title_run.font.size = Pt(28)
        title_run.bold = True
        title_run.font.name = "Arial"
        title_run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
        body = document.add_paragraph()
        body_run = body.add_run("Testo personalizzato")
        body_run.font.size = Pt(9)
        body_run.italic = True

        with tempfile.TemporaryDirectory() as temporary_directory:
            workspace = Path(temporary_directory)
            source = workspace / "styles.docx"
            document.save(source)
            html, css = pdf4um_app._docx_to_html(source)
            output = pdf4um_app._document_to_pdf(source, workspace)

            page = PdfReader(output).pages[0]
            styled_runs = {}

            def record_run(text, _cm, _tm, font, font_size):
                if text.strip():
                    styled_runs[text.strip()] = (font_size, str(font.get("/BaseFont")))

            page.extract_text(visitor_text=record_run)

        self.assertIn('font-family: "Helvetica"', css)
        self.assertIn("font-size: 28pt", css)
        self.assertIn("color: #123456", css)
        self.assertAlmostEqual(styled_runs["Titolo personalizzato"][0], 28)
        self.assertIn("Bold", styled_runs["Titolo personalizzato"][1])
        self.assertAlmostEqual(styled_runs["Testo personalizzato"][0], 9)
        self.assertIn("Oblique", styled_runs["Testo personalizzato"][1])

    def test_docx_conversion_preserves_heading_style(self):
        document = pdf4um_app.Document()
        heading = document.add_heading("Heading style preserved", level=1)
        heading.style.font.size = Pt(26)

        with tempfile.TemporaryDirectory() as temporary_directory:
            workspace = Path(temporary_directory)
            source = workspace / "heading.docx"
            document.save(source)
            html, css = pdf4um_app._docx_to_html(source)
            output = pdf4um_app._document_to_pdf(source, workspace)

            text_runs = []
            PdfReader(output).pages[0].extract_text(
                visitor_text=lambda text, _cm, _tm, _font, size:
                    text_runs.append((text, size)) if text.strip() else None
            )

        self.assertIn('<h1 class="pdf4um-paragraph-0"', html)
        self.assertIn("font-size: 26pt", css)
        self.assertTrue(any("Heading style preserved" in text and size == 26 for text, size in text_runs))

    def test_docx_conversion_preserves_parenthesized_list_markers(self):
        html = pdf4um_app._DocxListMarkerFormatter([1]).format(
            "<ol><li>Primo elemento</li><li>Secondo elemento</li></ol>"
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "numbering.pdf"
            pdf4um_app._html_to_pdf(
                html,
                ".pdf4um-custom-numbering { list-style-type: none; }",
                output,
            )
            text = PdfReader(output).pages[0].extract_text()

        self.assertIn("1) Primo elemento", text)
        self.assertIn("2) Secondo elemento", text)

    def test_odt_conversion_preserves_numbered_lists_and_grouped_table_rows(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            source = Path(temporary_directory) / "Report.odt"
            source.write_bytes(make_odt())

            html, _ = pdf4um_app._odt_to_html(source)

        response = self.post(
            "/api/documents-to-pdf",
            {"file": self.upload("Report.odt", make_odt())},
        )

        self.assertIn("<ol>", html)
        self.assertIn("</ol>", html)
        self.assertEqual(html.count("<tr>"), 2)
        self.assertIn("Prima colonna", html)
        self.assertIn("Seconda colonna", html)
        self.assertIn("Terza colonna", html)
        self.assertIn("Quarta colonna", html)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/pdf")
        extracted_text = "\n".join(
            page.extract_text() or "" for page in PdfReader(io.BytesIO(response.data)).pages
        )
        self.assertIn("Elemento elenco", extracted_text)
        self.assertIn("Seconda colonna", extracted_text)
        self.assertIn("Quarta colonna", extracted_text)

    def test_document_conversion_rejects_unsupported_and_invalid_files(self):
        unsupported = self.post(
            "/api/documents-to-pdf",
            {"file": self.upload("report.pdf", self.pdf)},
        )
        self.assertEqual(unsupported.status_code, 415)

        invalid = self.post(
            "/api/documents-to-pdf",
            {"file": self.upload("report.docx", b"not a DOCX archive")},
        )
        self.assertEqual(invalid.status_code, 422)

    def test_rename_sequential_and_random(self):
        sequential = self.post(
            "/api/rename",
            {
                "files": [self.upload("First.pdf", self.pdf), self.upload("Second.pdf", self.pdf)],
                "mode": "sequential",
                "prefix": "Chapter-",
                "startingNumber": "4",
            },
        )
        self.assertEqual(sequential.status_code, 200)
        self.assertEqual(self.zip_names(sequential), ["Chapter-4.pdf", "Chapter-5.pdf"])

        random_result = self.post(
            "/api/rename",
            {
                "files": self.upload("First.pdf", self.pdf),
                "mode": "random",
                "codeLength": "8",
                "useSymbols": "false",
            },
        )
        self.assertEqual(random_result.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(random_result.data)) as archive:
            self.assertIn("esporta.csv", archive.namelist())
            self.assertRegex(archive.read("esporta.csv").decode(), r"(?s)\.pdf;\d{8}\n")

    def test_encryption_and_text_download(self):
        encrypted = self.post(
            "/api/encrypt",
            {
                "files": self.upload("Report.pdf", self.pdf),
                "password": "test-password",
                "allowOpen": "false",
            },
        )
        self.assertEqual(encrypted.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(encrypted.data)) as archive:
            reader = PdfReader(io.BytesIO(archive.read("Report_cript.pdf")))
            self.assertTrue(reader.is_encrypted)
            self.assertTrue(reader.decrypt("test-password"))

        extracted = self.post(
            "/api/extract-text",
            {"file": self.upload("Notes.pdf", make_text_pdf("EXTRACTED TEXT"))},
        )
        self.assertEqual(extracted.status_code, 200)
        self.assertEqual(extracted.mimetype, "text/plain")
        self.assertIn("Notes_text.txt", extracted.headers["Content-Disposition"])
        self.assertIn(b"EXTRACTED TEXT", extracted.data)

    def test_rejects_invalid_uploads_and_unknown_routes_as_json(self):
        too_few = self.post(
            "/api/merge",
            {"files": self.upload("Report.pdf", self.pdf)},
        )
        self.assertEqual(too_few.status_code, 400)
        self.assertIn("error", json.loads(too_few.data))

        unsupported = self.post(
            "/api/split",
            {"file": self.upload("image.png", b"not a PDF")},
        )
        self.assertEqual(unsupported.status_code, 415)

        unreadable = self.post(
            "/api/extract-text",
            {"file": self.upload("broken.pdf", b"not a PDF")},
        )
        self.assertEqual(unreadable.status_code, 422)

        unknown = self.get("/api/unknown")
        self.assertEqual(unknown.status_code, 404)
        self.assertEqual(unknown.mimetype, "application/json")


if __name__ == "__main__":
    unittest.main()
